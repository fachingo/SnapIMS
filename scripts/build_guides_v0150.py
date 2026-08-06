#!/usr/bin/env python3
from __future__ import annotations
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
EVIDENCE = ROOT / "browser-evidence" / "v0.15.0"
VERSION = "0.15.0"
SCHEMA = "16 / 16"
DATE = "August 5, 2026"


def shade(cell, fill="E8EEF7"):
    props = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), fill)
    props.append(node)


def page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.add_run("Page ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def new_doc(title: str, subtitle: str) -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.65)
    sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.75)
    sec.right_margin = Inches(0.75)
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(10)
    for name, size in (("Title", 24), ("Heading 1", 17), ("Heading 2", 13), ("Heading 3", 11)):
        doc.styles[name].font.name = "Arial"
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].font.bold = True
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(title)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(subtitle)
    r.bold = True
    r.font.size = Pt(13)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"Active release v{VERSION} · inventory schema {SCHEMA} · {DATE}")
    doc.add_paragraph(
        "This guide describes the verified SnapIMS browser interface and operator workflow. Shopify draft creation is the default; live publication requires an explicit typed confirmation.",
        style="Intense Quote",
    )
    doc.add_page_break()
    for section in doc.sections:
        section.header.paragraphs[0].text = f"SnapIMS v{VERSION} · {title}"
        section.header.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        page_number(section.footer.paragraphs[0])
    return doc


def h(doc, text, level=1):
    doc.add_heading(text, level=level)


def para(doc, text):
    doc.add_paragraph(text)


def bullets(doc, values):
    for value in values:
        doc.add_paragraph(value, style="List Bullet")


def numbered(doc, values):
    for value in values:
        doc.add_paragraph(value, style="List Number")


def code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.name = "Courier New"
    r.font.size = Pt(9)


def table(doc, headers, rows):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.style = "Table Grid"
    for index, header in enumerate(headers):
        tbl.rows[0].cells[index].text = str(header)
        shade(tbl.rows[0].cells[index])
        for run in tbl.rows[0].cells[index].paragraphs[0].runs:
            run.bold = True
    for row in rows:
        cells = tbl.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = str(value)
    return tbl


def image(doc, filename, caption):
    path = EVIDENCE / filename
    if not path.exists():
        return
    with Image.open(path) as im:
        width_px, height_px = im.size
    ratio = width_px / max(1, height_px)
    max_w, max_h = 6.6, 6.6
    width = min(max_w, max_h * ratio)
    height = width / ratio
    doc.add_picture(str(path), width=Inches(width), height=Inches(height))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(caption)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].italic = True


def save(doc, filename):
    DOCS.mkdir(parents=True, exist_ok=True)
    out = DOCS / filename
    doc.save(out)
    print(out)


def add_shopify_workflow(doc, detailed=False):
    h(doc, "Shopify publish workflow")
    numbered(doc, [
        "Open Publish and select the working batch.",
        "Resolve all validation blockers. Live actions remain unavailable while blockers exist.",
        "Run Shopify Simulation. Simulation performs no Shopify write.",
        "Choose Create Shopify Drafts. SnapIMS uploads only publish-ready items and records Product, Variant, and Inventory Item GIDs.",
        "After success, SnapIMS opens Shopify draft products in a separate browser tab. The Publish page also provides Open in Shopify, Copy Product Link, and Copy Product GID controls.",
        "Review drafts when desired, then type SUBMIT and choose Publish Drafts Live.",
        "For deliberate direct publication, type SUBMIT LIVE. This bypass is not the default path.",
    ])
    if detailed:
        table(doc, ["Action", "Required confirmation", "Effect"], [
            ("Create Shopify Drafts", "None after successful simulation", "Creates unpublished Shopify products"),
            ("Publish Drafts Live", "SUBMIT", "Activates products and publishes them to the selected publication"),
            ("Publish Directly", "SUBMIT LIVE", "Creates drafts and then publishes them live"),
            ("Archive", "ARCHIVE", "Makes selected Shopify products inactive"),
            ("Restore Draft", "RESTORE", "Returns selected products to draft state"),
            ("Delete", "DELETE", "Permanently deletes selected Shopify products"),
            ("Keep Shopify", "KEEP SHOPIFY", "Imports the remote product snapshot into supported SnapIMS fields"),
            ("Merge", "MERGE", "Applies the supported conflict merge operation"),
        ])
        bullets(doc, [
            "Use checkboxes for selected-item operations; use the explicit entire-batch scope only when intended.",
            "Select All, Invert Selection, search, filtering, sorting, retry, reconciliation, archive, restore, and delete are available from the management section.",
            "Durable jobs show pending, running, retrying, complete, skipped, and failed rows with progress, elapsed time, ETA, and failure details.",
            "Restarting SnapIMS recovers unfinished jobs. Completed rows are not repeated.",
            "SKU reconciliation and stored Shopify GIDs prevent accidental duplicate creation.",
        ])


def day_to_day():
    doc = new_doc("SnapIMS Day-to-Day Guide", "First-time operator workflow")
    h(doc, "1. Start and verify")
    code(doc, 'export PATH="$HOME/.local/bin:$PATH"\nsnapims up\nsnapims status')
    bullets(doc, ["Version must show 0.15.0.", f"Inventory schema must show {SCHEMA}.", "Open http://127.0.0.1:8767/."])
    h(doc, "2. Import one folder")
    numbered(doc, [
        "Create one folder inside ~/SnapIMS-data/batches. The folder name becomes the default batch name.",
        "Put photographs inside in capture order. Use only the NEXT ITEM QR between tapes.",
        "Open Import, select the folder, enter a free-text location or leave it unassigned, then Preview.",
        "Correct grouping if needed and Commit once the item count is correct.",
    ])
    image(doc, "chromium-import-preview-summary.png", "Import Preview with durable progress and interpreted items.")
    h(doc, "3. Recognize, review, and edit")
    bullets(doc, [
        "Run recognition when configured, then review every product field and evidence.",
        "Use Batch Editor for paging, search, filtering, exact Fill Down, Append Tags, and audited bulk changes.",
        "Hidden rows are deselected automatically so bulk actions cannot silently affect filtered-out records.",
    ])
    image(doc, "chromium-batch-editor.png", "Batch Editor with bounded paging and selection-safe bulk tools.")
    h(doc, "4. Connect Shopify")
    numbered(doc, [
        "Open Settings → Shopify connection.",
        "Enter the permanent .myshopify.com domain, Client ID, Client Secret, and current SnapIMS administrator password.",
        "Test and save securely, then confirm store identity, scopes, token status, and location discovery.",
        "Select the inventory location by name and the publication used for live products.",
        "Keep draft-first publishing enabled.",
    ])
    image(doc, "chromium-shopify-settings.png", "Verified Shopify client-credentials connection and named location/publication setup.")
    h(doc, "5. Publish")
    add_shopify_workflow(doc)
    image(doc, "chromium-publish-workflow.png", "Publish page: validation, simulation, draft creation, live publication, and product management.")
    image(doc, "chromium-publish-drafts-complete.png", "Completed draft job with Shopify links and durable item results.")
    h(doc, "6. End of batch")
    numbered(doc, [
        "Confirm physical count, titles, prices, tags, condition, quantity, and location.",
        "Verify draft/live results and failed rows independently.",
        "Resolve conflicts or retry failures before archiving the batch.",
        "Stop SnapIMS before an offline backup.",
    ])
    code(doc, "snapims down")
    save(doc, f"SnapIMS_v{VERSION}_Day-to-Day_Guide.docx")


def operator():
    doc = new_doc("SnapIMS Operator Guide", "Complete verified operator procedures")
    h(doc, "1. Release identity")
    table(doc, ["Identity", "Expected"], [("Application", VERSION), ("Inventory schema", SCHEMA), ("Data directory", "~/SnapIMS-data"), ("Batch Home", "~/SnapIMS-data/batches")])
    h(doc, "2. Install or upgrade")
    code(doc, "unzip SnapIMS-v0.15.0-full-source.zip\ncd SnapIMS-v0.15.0\nchmod +x install_v0150.sh\nSNAPIMS_DATA_DIR=\"$HOME/SnapIMS-data\" ./install_v0150.sh\nsnapims status\nsnapims up")
    para(doc, "Back up the data directory before upgrade. The installer and runtime reject unsupported schema/application combinations.")
    h(doc, "3. Import protocol")
    para(doc, "Folder boundaries define the batch. NEXT ITEM is the only routine QR command. Batch name and location are entered or overridden in the browser.")
    table(doc, ["Situation", "Result"], [
        ("No NEXT card", "Valid one-item batch"), ("Leading/trailing NEXT", "Ignored with warning"), ("Consecutive NEXT", "Collapsed with warning"), ("Consumer QR", "Product photo"), ("Legacy START/END/location QR", "Ignored legacy command"),
    ])
    h(doc, "4. Review and Batch Editor")
    bullets(doc, [
        "Review immutable Item IDs, photos, title, price, description, tags, quantity, condition, location, Rare, and Physical Review.",
        "Fill Down supports Tags, Location, Review, Rare, Title, Price, Discount, and Description.",
        "Tags Fill Down replaces the destination set; Append Tags adds without replacing.",
        "Stale edits are rejected rather than overwriting newer changes.",
    ])
    image(doc, "chromium-batch-editor-controls.png", "Batch Editor controls for selection, filtering, and exact bulk changes.")
    h(doc, "5. Shopify setup")
    bullets(doc, [
        "Required capabilities are derived from actual operations: locations, products, inventory, media, and publications.",
        "SnapIMS stores Client Secret and access-token cache encrypted and refreshes the token automatically.",
        "Select the active inventory location and the live publication by name.",
        "The saved legacy static token is deprecated and should be removed after client credentials pass.",
    ])
    image(doc, "chromium-shopify-section.png", "Shopify settings with read-only verification and recovery controls.")
    add_shopify_workflow(doc, detailed=True)
    image(doc, "chromium-publish-workflow.png", "Full v0.15.0 publish workflow.")
    h(doc, "6. Conflict and product management")
    bullets(doc, [
        "Reconcile compares supported local and remote values and records a conflict when they differ.",
        "Keep SnapIMS pushes the supported local state to Shopify.",
        "Keep Shopify imports the supported remote state into SnapIMS.",
        "Archive is reversible; Delete is permanent and requires DELETE.",
        "Single-item, selected-item, and entire-batch scopes are explicit.",
    ])
    h(doc, "7. Recovery")
    table(doc, ["Condition", "Action"], [
        ("Interrupted Shopify job", "Restart SnapIMS, reopen Publish, and resume/retry unfinished rows."),
        ("Expired token", "No operator action normally; token refresh is automatic."),
        ("Missing publication", "Return to Settings and select a publication before live publish."),
        ("Shopify rate limit/transport failure", "Wait and retry failed rows; completed rows are skipped."),
        ("Wrong version/path", "Run snapims status and reinstall from the intended extracted release."),
        ("Database problem", "Stop SnapIMS, preserve the data directory, and restore only from a copied backup."),
    ])
    h(doc, "8. Diagnostics and CLI")
    code(doc, "snapims status\nsnapims doctor\nsnapims logs\nsnapims shopify status\nsnapims shopify test\nsnapims shopify refresh\nsnapims shopify jobs")
    h(doc, "9. Backup")
    code(doc, 'snapims down\ncp -a "$HOME/SnapIMS-data" "$HOME/SnapIMS-data-backup-$(date +%Y%m%d-%H%M%S)"')
    save(doc, f"SnapIMS_v{VERSION}_Operator_Guide.docx")


def manual():
    doc = new_doc("SnapIMS Complete Operating Manual", "Architecture, operations, recovery, and release acceptance")
    h(doc, "1. Operating model")
    table(doc, ["Layer", "Purpose"], [
        ("Source photos", "Immutable operator originals"), ("Import cache", "Restart-safe decoding, QR interpretation, and thumbnails"), ("Inventory database", "Batches, items, controlled tags, history, and checkpoints"), ("Recognition", "Durable provider attempts and evidence"), ("Review/Batch Editor", "Human decisions and audited edits"), ("Shopify jobs", "Durable draft, live, sync, reconcile, archive, restore, and delete operations"),
    ])
    h(doc, "2. Shopify state model")
    table(doc, ["State", "Meaning"], [
        ("Draft", "Created in Shopify but not live"), ("Published", "Active and published to the selected publication"), ("Archived", "Inactive on Shopify"), ("Deleted", "Remote product permanently removed"), ("Needs Sync", "Local/remote state requires action"), ("Conflict", "Supported local and remote values differ"),
    ])
    add_shopify_workflow(doc, detailed=True)
    image(doc, "chromium-publish-workflow.png", "Desktop publish workflow browser evidence.")
    image(doc, "chromium-mobile-publish.png", "Publish workflow verified at a 390-pixel mobile viewport.")
    h(doc, "3. Durable jobs and idempotency")
    bullets(doc, [
        "Jobs and per-item stages are stored in schema 16.",
        "A completed item is not executed again after refresh or restart.",
        "Concurrent duplicate work is rejected or resolves to the existing durable job.",
        "Transient failures use bounded retry; authentication failure permits one token refresh and retry.",
        "SKU reconciliation and persisted Shopify GIDs prevent duplicate product creation.",
    ])
    h(doc, "4. Security")
    bullets(doc, [
        "Client Secret and access-token cache are encrypted and owner-only.",
        "Secrets are not written to logs, HTML, screenshots, browser storage, or audit details.",
        "All state-changing browser routes remain authenticated and CSRF-protected.",
        "Destructive or public-storefront actions require exact typed confirmation.",
        "Connection testing is read-only; live writes occur only from the Publish workflow.",
    ])
    h(doc, "5. Browser and automated verification")
    table(doc, ["Evidence", "Result"], [
        ("Automated suite", "299 passed, 3 skipped, 0 failed"),
        ("Browser checks", "52 / 52 passed"),
        ("Console/page errors", "0"),
        ("Unexpected failed requests", "0"),
        ("HTTP 500 responses", "0"),
        ("Desktop Chromium", "Pass through local relay"),
        ("390 px mobile", "Pass"),
        ("Upgrade from v0.14.0", "Pass; schema 16; inventory preserved"),
    ])
    h(doc, "6. External tests not claimed")
    table(doc, ["Capability", "Reason"], [
        ("Live owner-store product writes", "No owner Shopify credentials were available in the sandbox; deterministic fakes covered mutations and state transitions."),
        ("Live OpenAI", "No production API key in sandbox."),
        ("Native Firefox", "Runtime unavailable."),
        ("Target 2012 Mac mini", "Hardware unavailable."),
        ("Owner Cloudflare/Guacamole/xrdp", "Infrastructure unavailable."),
        ("Real HEIC", "HEIC runtime unavailable."),
    ])
    h(doc, "7. v1.0 acceptance")
    bullets(doc, [
        "Complete the real 20-tape Pixel pilot.",
        "Create and inspect real Shopify drafts, then publish a controlled live product and test archive/delete recovery decisions.",
        "Run live AI, real CSV round trip, restart durability, backup/restore, target-hardware soak, direct Firefox/Chromium walkthrough, and security review.",
        "Follow the final Operator Guide against the production browser UI and resolve every discrepancy.",
        "Do not label 1.0.0 while any Critical or High production blocker remains.",
    ])
    h(doc, "8. Evidence locations")
    bullets(doc, ["browser-evidence/v0.15.0/", "release-evidence/v0.15.0/", "TEST_RESULTS.md", "BROWSER_VERIFICATION.md", "ISSUE_TRACEABILITY_V0150.md", "V1_PRODUCTION_ACCEPTANCE_GAPS.md"])
    save(doc, f"SnapIMS_v{VERSION}_Complete_Operating_Manual.docx")


if __name__ == "__main__":
    day_to_day()
    operator()
    manual()
