from __future__ import annotations

from pathlib import Path
from typing import Iterable

from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SHOTS = ROOT / "operator-audit-assets" / "v0.6.0-browser" / "screenshots"
SLMC_SHOTS = ROOT / "operator-audit-assets" / "slmc-0.1.0-browser" / "screenshots"
TMP = ROOT / ".guide-assets-slmc"
OUT = ROOT / "SnapIMS_Operator_Guide.docx"
VERSION = "0.6.0 + SLMC-0.1.0"

ACCENT = "8055D9"
DARK = "211C2D"
MUTED = "6A6475"
LIGHT = "F4F0FA"
GREEN = "2F7D5A"
RED = "A9364D"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("SnapIMS 0.6.0 + SLMC-0.1.0 Operator Guide  |  ")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, text, fld_end])


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Liberation Sans"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for name, size, color in (
        ("Title", 30, DARK),
        ("Subtitle", 13, MUTED),
        ("Heading 1", 20, ACCENT),
        ("Heading 2", 14, DARK),
        ("Heading 3", 11, ACCENT),
    ):
        style = styles[name]
        style.font.name = "Liberation Sans"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = name != "Subtitle"


def configure_sections(doc: Document) -> None:
    for section in doc.sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.55)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)
        section.header_distance = Inches(0.2)
        section.footer_distance = Inches(0.25)
        add_page_number(section.footer.paragraphs[0])


def add_callout(doc: Document, title: str, body: str, kind: str = "info") -> None:
    table = doc.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.cell(0, 0)
    fill = {"info": LIGHT, "success": "EAF6F0", "warning": "FFF4DB", "danger": "FCE8EC"}[kind]
    set_cell_shading(cell, fill)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string({"info": ACCENT, "success": GREEN, "warning": "8A5A00", "danger": RED}[kind])
    p2 = cell.add_paragraph(body)
    p2.paragraph_format.space_after = Pt(0)
    doc.add_paragraph()


def add_steps(doc: Document, steps: Iterable[str]) -> None:
    for i, step in enumerate(steps, 1):
        p = doc.add_paragraph()
        p.style = doc.styles["Normal"]
        p.paragraph_format.left_indent = Inches(0.12)
        p.paragraph_format.first_line_indent = Inches(-0.12)
        r = p.add_run(f"{i}. ")
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(ACCENT)
        p.add_run(step)


def screenshot_slices(path: Path, max_height: int = 920) -> list[Path]:
    TMP.mkdir(exist_ok=True)
    image = Image.open(path).convert("RGB")
    outputs: list[Path] = []
    if image.height <= max_height:
        output = TMP / path.name
        image.save(output, quality=95)
        return [output]

    # Split tall browser captures into balanced, overlapping views. Balanced
    # slices avoid a final page containing only a narrow strip of the UI.
    import math

    part_count = math.ceil(image.height / max_height)
    slice_height = math.ceil(image.height / part_count)
    overlap = 45
    for idx in range(part_count):
        top = max(0, idx * slice_height - (overlap if idx else 0))
        bottom = min(image.height, (idx + 1) * slice_height + (overlap if idx < part_count - 1 else 0))
        crop = image.crop((0, top, image.width, bottom))
        output = TMP / f"{path.stem}-part-{idx + 1:02d}.png"
        crop.save(output)
        outputs.append(output)
    return outputs


def add_screenshot(doc: Document, filename: str, caption: str) -> None:
    if filename.startswith("slmc:"):
        path = SLMC_SHOTS / filename.removeprefix("slmc:")
    else:
        path = SHOTS / filename
    if not path.exists():
        raise FileNotFoundError(path)
    parts = screenshot_slices(path)
    for idx, part in enumerate(parts, 1):
        if idx > 1:
            doc.add_page_break()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run()
        run.add_picture(str(part), width=Inches(7.0))
        cp = doc.add_paragraph(caption if len(parts) == 1 else f"{caption} - view {idx} of {len(parts)}")
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_after = Pt(8)
        for r in cp.runs:
            r.italic = True
            r.font.size = Pt(8)
            r.font.color.rgb = RGBColor.from_string(MUTED)


def section(doc: Document, title: str, intro: str, steps: list[str], screenshot: str | None = None, caption: str = "") -> None:
    doc.add_heading(title, level=1)
    doc.add_paragraph(intro)
    add_steps(doc, steps)
    if screenshot:
        add_screenshot(doc, screenshot, caption)
    doc.add_page_break()


def build() -> None:
    doc = Document()
    configure_styles(doc)
    configure_sections(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(60)
    r = title.add_run("SnapIMS")
    r.bold = True
    r.font.name = "Liberation Sans"
    r.font.size = Pt(42)
    r.font.color.rgb = RGBColor.from_string(ACCENT)
    p = doc.add_paragraph("Operator Guide")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.style = doc.styles["Title"]
    p2 = doc.add_paragraph("Application 0.6.0 · SLMC-0.1.0 Integration Build")
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.style = doc.styles["Subtitle"]
    doc.add_paragraph()
    add_callout(
        doc,
        "Purpose",
        "SnapIMS converts an ordered photo stream into durable inventory records. The SLMC integration adds a permanent local Movie catalog, local-first matching, bounded Wikipedia discovery on genuine misses, and structured CSV/Shopify simulation fields without slowing the one-action Review workflow.",
        "info",
    )
    add_callout(
        doc,
        "Integration status",
        "The application remains SnapIMS 0.6.0. SLMC-0.1.0 is an isolated integration package, not a final SnapIMS release and not 1.0. Live AI, live Wikipedia, a real Pixel batch, one live Shopify draft, physical CSV reconciliation, and final independent acceptance remain required.",
        "warning",
    )
    doc.add_page_break()

    doc.add_heading("Fast Start", level=1)
    doc.add_paragraph("Launch SnapIMS from the repository virtual environment:")
    code = doc.add_paragraph()
    code.style = doc.styles["Normal"]
    run = code.add_run("cd ~/Projects/SnapIMS-v0.6.0\nsource .venv/bin/activate\nsnapims --data-dir ~/SnapIMS-data serve")
    run.font.name = "Liberation Mono"
    run.font.size = Pt(9)
    add_callout(doc, "Browser address", "Open http://127.0.0.1:8767. Keep the terminal running. Ctrl+C stops the server.", "success")
    add_steps(doc, [
        "Import a folder containing the chronological camera photo stream.",
        "Run recognition. After each successful recognition result is committed, SnapIMS searches the permanent local Movie catalog before any Wikipedia request.",
        "Resolve items one at a time in Review. Normal unique catalog matches require no extra confirmation click.",
        "Use CSV download/upload only when an external spreadsheet is useful.",
        "Validate the current working batch in Publish, then simulate or create the requested output.",
    ])
    doc.add_page_break()

    section(doc, "1. Home Dashboard", "Home shows the active batches, recognition state, unfinished work, and the next logical actions.", [
        "Choose Import to create a new batch.",
        "Open Review for the one-tape workstation.",
        "Open Batch Editor for bulk changes and low-confidence triage.",
        "Open Publish only after the working batch has no blockers.",
    ], "01-home-dashboard.png", "Home dashboard and batch actions")

    section(doc, "2. Import a Photo Folder", "Normal import no longer requires repeatedly typing a Linux path. Browse Folder is the preferred path; configured and recent folders are one-click shortcuts.", [
        "Click Browse Folder and select the flat camera-export folder.",
        "If the native chooser is unavailable, open Advanced Manual Path.",
        "Click Preview. Preview does not create durable item identity.",
        "Confirm the image count, QR events, item boundaries, and warnings.",
        "Click Preserve and Import Batch once. Duplicate submission protection prevents accidental double import.",
    ], "03-native-folder-selected.png", "Native folder selection on Import")

    section(doc, "3. Preview and Commit", "The preview is the last safe checkpoint before SnapIMS preserves originals, creates durable identities, and generates media derivatives.", [
        "Verify BATCH START and BATCH END were detected.",
        "Verify every ITEM NEXT boundary created one intended tape.",
        "Inspect warning messages before committing.",
        "Commit the batch. SnapIMS preserves originals and generates fast browser and recognition copies.",
    ], "04-import-preview.png", "Import preview with QR and item summary")

    section(doc, "4. Recognition and Failure Recovery", "Recognition can use Mock for testing or a configured provider for live work. A provider failure never leaves a dead button or silent refresh.", [
        "Start identification from Review or the batch action.",
        "When recognition fails, open Review Failure.",
        "Read the human-readable reason, provider, timestamp, completed count, and remaining count.",
        "Choose Retry Failed Items, another provider, Manual Review, Skip Recognition, or Diagnostics.",
        "Manual Review leaves items unfinished and editable; it does not falsely mark them recognized.",
    ], "06-recognition-failure-recovery.png", "Recognition failure recovery workspace")

    section(doc, "5. Fast Individual Review", "Review is optimized for the routine loop: look, confirm, optionally type a price, press Enter, and continue.", [
        "Confirm the displayed title. SnapIMS uses manual, saved, CSV, then AI values in that order.",
        "The Price field is focused and selected when an unfinished tape opens.",
        "Type a replacement price or leave the current value unchanged.",
        "Press Enter to run Approve & Next. The next unfinished tape opens with Price selected again.",
        "Use Later to postpone the item without marking it complete.",
        "Use full Edit only for deeper corrections or a genuinely missing title.",
    ], "08-price-enter-next.png", "Price -> Enter -> next Price workflow")

    section(doc, "6. Catalog Match States", "The compact catalog strip is evidence, not a second approval workflow. The photograph remains authority and Approve & Next remains the routine action.", [
        "Local Match means the recognized title reused an existing permanent Movie record and made no Wikipedia request.",
        "New Catalog Record means a bounded Wikipedia miss flow created one permanent Movie record with source provenance.",
        "Searching means the background external lookup is still running; Review remains usable.",
        "Ambiguous means two or more credible films remain. Open the candidate panel or Edit details before approval.",
        "Not Found, Failed, Needs Review, or Catalog Unavailable leave the original AI suggestion visible and do not destroy the physical Item.",
        "The displayed title precedence is operator-saved title, unique catalog title, AI suggestion, then manual Untitled state.",
    ], "slmc:03-local-match-review.png", "Compact local catalog match in the normal Review workstation")

    section(doc, "7. First Miss and Permanent Reuse", "Wikipedia is an external discovery source only. SnapIMS stores normalized Movie facts permanently in movie_catalog.sqlite3 and searches that database first on every future recognition.", [
        "A genuine local miss queues a bounded English Wikipedia lookup through the official MediaWiki endpoint.",
        "SnapIMS normally inspects no more than five candidates and rejects obvious songs, albums, books, people, companies, and unrelated episodes.",
        "A materially unique film creates one permanent Movie, aliases, source provenance, and an Item link.",
        "A later copy of the same movie reuses that Movie ID with zero new Wikipedia request.",
        "Wikipedia can establish film-level facts; it cannot establish exact VHS distributor, UPC, Canadian edition, cover variant, condition, shelf, price, discount, or quantity.",
    ], "slmc:04-new-wikipedia-record-and-next.png", "New catalog record created after one bounded Wikipedia fixture lookup")

    section(doc, "8. Repeated Copies", "Repeated tapes remain separate physical Items even when they share one Movie record.", [
        "Confirm that each tape keeps its own immutable Item ID, Batch ID, images, shelf, condition, price, discount, quantity, Review history, and Shopify state.",
        "The shared Movie ID supplies reusable canonical title and factual metadata only.",
        "Do not pool copies or infer a shared VHS Edition from the Movie match.",
        "A corrected title marks the prior link stale, preserves link history, and starts a new local-first lookup.",
    ], "slmc:06-second-copy-local-reuse.png", "Second physical copy reusing the permanent local Movie record")

    section(doc, "9. Batch Editor", "Batch Editor is the bulk operator workstation. It edits the same authoritative working records used by Review and Publish.", [
        "Use search, status filters, confidence threshold, confidence buckets, and sorting to expose exceptions.",
        "Edit Title, Price, Discount, Description, Location, flags, and other working fields directly.",
        "Watch the save indicator. Failed edits remain visible instead of disappearing.",
        "Use Ctrl+S to save pending work, Ctrl+F to focus search, Ctrl+A to select visible rows, and Ctrl+Z/Ctrl+Y to undo or redo local edits.",
        "Click a thumbnail for the fast media viewer or confidence for stored recognition evidence.",
    ], "09-batch-editor-bulk-price.png", "Batch Editor with health, confidence, and bulk controls")

    section(doc, "10. Command Palette and Bulk Operations", "The command palette supplements visible controls and keeps frequent actions keyboard-accessible.", [
        "Press Ctrl+Shift+P to open the command palette.",
        "Jump to Review, Batch Editor, Publish, Import, Diagnostics, CSV tools, or batch rollback.",
        "Select rows and preview bulk operations before applying them.",
        "For ToonieTape work, select the intended rows, choose Bulk Price, enter $4.00, inspect the affected count, then Apply.",
        "SnapIMS creates a checkpoint before large destructive operations.",
    ], "10-command-palette.png", "Keyboard command palette")

    section(doc, "11. CSV Round Trip and Difference Preview", "CSV remains available for external spreadsheet work, but uploaded values are never applied blindly.", [
        "Download the current working CSV from Publish or Batch Tools.",
        "Edit the file in Excel or LibreOffice without changing immutable Item IDs.",
        "Upload the edited CSV. SnapIMS stages it instead of modifying the batch immediately.",
        "Use the Difference Preview to inspect matched rows, missing items, validation errors, and field-level changes.",
        "Cancel to prove no values changed, or confirm Apply Valid Changes after all blocking errors are resolved.",
        "Use Undo CSV Import to restore the durable checkpoint.",
    ], "11-csv-difference-preview.png", "CSV difference preview before application")

    section(doc, "12. External Review", "A valid CSV can satisfy review for large intentionally uniform batches, but only after explicit operator confirmation.", [
        "Apply the validated CSV changes first.",
        "Confirm that all required titles exist, SKUs are unique, and no blocking row remains.",
        "Choose Mark Uploaded Batch as Externally Reviewed.",
        "Accept the responsibility statement. SnapIMS records CSV_EXTERNAL_REVIEW as the review source.",
        "Rows with unresolved errors remain unfinished.",
    ], "12-csv-applied.png", "Applied CSV values in the authoritative working batch")

    section(doc, "13. Publish", "Publish validates the authoritative working batch regardless of whether edits came from AI, Review, Batch Editor, or CSV.", [
        "Review all publish blockers and follow links back to the affected item or editor row.",
        "Use Shopify simulation before any live publishing test.",
        "Download the final CSV and reconcile it against the physical batch.",
        "Live publishing must create drafts, use idempotency checkpoints, and never duplicate a product on retry.",
        "A live Shopify draft remains a 1.0 acceptance requirement.",
    ], "13-publish-simulation.png", "Publish validation and Shopify simulation")

    section(doc, "14. Diagnostics and Catalog Recovery", "Diagnostics reports inventory health and the independent permanent Movie catalog without leaking credentials.", [
        "Confirm Inventory integrity is ok, Inventory FK violations are zero, and Inventory schema is 7 for this integration build.",
        "Confirm Catalog integrity is ok, Catalog FK violations are zero, and Catalog schema is 1.",
        "Review Movie, alias, linked-item, local-hit, Wikipedia-call, ambiguous, failed-job, stale-link, database-size, and FTS values.",
        "Use Back up catalog before migrations, bulk imports, merge, split, restore, or server movement.",
        "Use Rebuild search index only as a maintenance action; it does not rebuild the authoritative Movie data.",
        "Retry failed lookups after correcting connectivity or provider errors. Recognition history and physical Items remain intact.",
        "If movie_catalog.sqlite3 is unavailable or incompatible, do not delete it. Inventory, Review, Approve & Next, CSV, and Shopify simulation continue without pretending catalog metadata exists.",
    ], "slmc:08-catalog-diagnostics.png", "Independent inventory and permanent local movie-catalog diagnostics")

    doc.add_heading("15. End-of-Batch Checklist", level=1)
    add_steps(doc, [
        "Unfinished count is zero, or every remaining item is intentionally postponed and understood.",
        "No failed recognition state lacks an operator decision.",
        "Required titles, prices, quantities, locations, and SKUs pass validation.",
        "Low-confidence, ambiguous, failed, and flagged items were reviewed using stored evidence or the original photographs.",
        "Every Local Match or New Catalog Record agrees with the photograph; no same-title remake was silently accepted.",
        "The final CSV contains both immutable Item ID and the intended local Movie ID, while item-specific Price, Discount, condition, shelf, and quantity remain unchanged.",
        "Catalog integrity is ok, catalog foreign-key violations are zero, and a recent catalog backup exists.",
        "Bulk edits and CSV replacements were checkpointed and verified.",
        "Shopify simulation completed without duplicate or idempotency warnings.",
        "The final CSV was downloaded and reconciled against the physical batch.",
        "SnapIMS was restarted once and the current batch state remained correct.",
        "Do not label the project 1.0 until the real 20-tape Pixel pilot, live AI, live Shopify draft, CSV reconciliation, restart durability, final guide walkthrough, and browser verification all pass.",
    ])
    add_callout(doc, "Operator rule", "When SnapIMS reports an error, do not delete the database or originals. Preserve the traceback, copy Diagnostics, and recover from the last checkpoint or backup.", "danger")

    doc.core_properties.title = "SnapIMS Operator Guide"
    doc.core_properties.subject = "SnapIMS 0.6.0 with SLMC-0.1.0 integration workflows"
    doc.core_properties.author = "SnapIMS Project"
    doc.core_properties.keywords = "SnapIMS, SLMC, local movie catalog, Wikipedia, inventory, VHS, operator guide"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
