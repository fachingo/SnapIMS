from __future__ import annotations

import math
from pathlib import Path
from typing import Iterable

from PIL import Image
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SHOTS = ROOT / "release-evidence" / "v0.6.1" / "browser" / "screenshots"
ASSETS = ROOT / ".guide-assets-v061"
OUT = ROOT / "SnapIMS_Operator_Guide.docx"
VERSION = "0.6.1"

ACCENT = "8B5CF6"
DARK = "17131F"
MUTED = "676071"
LIGHT = "F3EEFC"
GREEN = "176B4D"
YELLOW = "8A5A00"
RED = "A9364D"
BORDER = "D9D2E6"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = BORDER) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), color)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(f"SnapIMS Operator Guide v{VERSION}  |  ")
    run.font.name = "Liberation Sans"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, text, fld_end])


def configure(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.52)
    section.bottom_margin = Inches(0.48)
    section.left_margin = Inches(0.67)
    section.right_margin = Inches(0.67)
    section.header_distance = Inches(0.15)
    section.footer_distance = Inches(0.20)
    add_page_number(section.footer.paragraphs[0])

    normal = doc.styles["Normal"]
    normal.font.name = "Liberation Sans"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.05

    styles = {
        "Title": (30, DARK),
        "Subtitle": (13, MUTED),
        "Heading 1": (19, ACCENT),
        "Heading 2": (13, DARK),
        "Heading 3": (10.5, ACCENT),
    }
    for name, (size, color) in styles.items():
        style = doc.styles[name]
        style.font.name = "Liberation Sans"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = name != "Subtitle"
        style.paragraph_format.space_after = Pt(6)


def add_callout(doc: Document, title: str, body: str, kind: str = "info") -> None:
    fills = {"info": LIGHT, "success": "E9F6F0", "warning": "FFF4D7", "danger": "FCE8EC"}
    accents = {"info": ACCENT, "success": GREEN, "warning": YELLOW, "danger": RED}
    table = doc.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_shading(cell, fills[kind])
    set_cell_border(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(accents[kind])
    p2 = cell.add_paragraph(body)
    p2.paragraph_format.space_after = Pt(0)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_steps(doc: Document, steps: Iterable[str]) -> None:
    for number, step in enumerate(steps, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.19)
        p.paragraph_format.first_line_indent = Inches(-0.19)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(f"{number}. ")
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(ACCENT)
        p.add_run(step)


def add_bullets(doc: Document, bullets: Iterable[str]) -> None:
    for text in bullets:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(text)


def screenshot_slices(path: Path, max_height: int = 1080) -> list[Path]:
    ASSETS.mkdir(exist_ok=True)
    image = Image.open(path).convert("RGB")
    if image.height <= max_height:
        out = ASSETS / f"guide-{path.name}"
        image.save(out, optimize=True)
        return [out]
    parts = math.ceil(image.height / max_height)
    base = math.ceil(image.height / parts)
    overlap = 65
    outputs: list[Path] = []
    for idx in range(parts):
        top = max(0, idx * base - (overlap if idx else 0))
        bottom = min(image.height, (idx + 1) * base + (overlap if idx < parts - 1 else 0))
        crop = image.crop((0, top, image.width, bottom))
        out = ASSETS / f"guide-{path.stem}-part-{idx + 1:02d}.png"
        crop.save(out, optimize=True)
        outputs.append(out)
    return outputs


def add_picture(doc: Document, path: Path, caption: str, width: float = 7.0) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Inches(width))
    cp = doc.add_paragraph(caption)
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(5)
    for run in cp.runs:
        run.italic = True
        run.font.size = Pt(7.5)
        run.font.color.rgb = RGBColor.from_string(MUTED)


def add_screenshot_pages(doc: Document, filename: str, caption: str) -> None:
    path = SHOTS / filename
    if not path.exists():
        raise FileNotFoundError(path)
    parts = screenshot_slices(path)
    for idx, part in enumerate(parts):
        if idx:
            doc.add_page_break()
            doc.add_heading(f"Screen continuation", level=2)
        add_picture(doc, part, caption if len(parts) == 1 else f"{caption} - view {idx + 1} of {len(parts)}")


def new_section(doc: Document, title: str, intro: str, steps: list[str], screenshot: str | None = None, caption: str = "") -> None:
    doc.add_heading(title, level=1)
    doc.add_paragraph(intro)
    add_steps(doc, steps)
    if screenshot:
        add_screenshot_pages(doc, screenshot, caption)
    doc.add_page_break()


def add_keyboard_table(doc: Document) -> None:
    rows = [
        ("Enter", "In the selected Price field, approve the current unfinished tape exactly once and open the next unfinished tape."),
        ("Ctrl+S", "Save pending Batch Editor edits."),
        ("Ctrl+F", "Focus the current search field."),
        ("Ctrl+Shift+P", "Open the command palette."),
        ("F2", "Edit the current Batch Editor cell where supported."),
        ("Escape", "Cancel the current dialog or edit when supported."),
        ("Tab / Shift+Tab", "Move through fields without leaving the keyboard."),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(1.4)
    table.columns[1].width = Inches(5.6)
    header = table.rows[0].cells
    header[0].text = "Shortcut"
    header[1].text = "Action"
    for cell in header:
        set_cell_shading(cell, ACCENT)
        set_cell_border(cell, ACCENT)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    for key, action in rows:
        cells = table.add_row().cells
        cells[0].text = key
        cells[1].text = action
        for cell in cells:
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cells[0].paragraphs[0].runs[0].bold = True


def build() -> None:
    doc = Document()
    configure(doc)

    cover = doc.add_paragraph()
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover.paragraph_format.space_before = Pt(40)
    r = cover.add_run("SnapIMS")
    r.font.name = "Liberation Sans"
    r.font.size = Pt(42)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(ACCENT)
    p = doc.add_paragraph("Operator Guide")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.style = doc.styles["Title"]
    p = doc.add_paragraph(f"Version {VERSION} - Mega Stabilization Patch")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.style = doc.styles["Subtitle"]
    add_callout(doc, "Operating model", "SnapIMS is an exception-handling workstation. Routine work is photo review, optional Price or Discount adjustment, then Approve & Next.", "info")
    add_callout(doc, "Release status", "Version 0.6.1 is not production 1.0. The real 20-tape Pixel pilot, live AI, one live Shopify draft, physical CSV verification, restart durability, and final blocker review still govern 1.0 acceptance.", "warning")
    add_picture(doc, SHOTS / "01-home-production.png", "SnapIMS 0.6.1 Home - browser-verified production mode", width=6.8)
    doc.add_page_break()

    doc.add_heading("Quick start and safety", level=1)
    doc.add_paragraph("The routine operator path is intentionally short:")
    add_steps(doc, [
        "Photograph START, shelf, each tape, NEXT between tapes, and END last.",
        "Import the complete camera folder and verify Preview before preserving anything.",
        "Identify with a live provider, or continue manually when recognition is blocked.",
        "Review the photograph, confirm the title, optionally adjust Price or Discount, then use Approve & Next.",
        "Use Batch Editor or CSV only for deliberate batch-wide changes.",
        "Run Shopify simulation before any live draft acceptance test.",
    ])
    add_callout(doc, "Identity rule", "Item ID is permanent. Never relink records by title, filename, spreadsheet row position, or a newly generated identifier.", "danger")
    add_callout(doc, "Source of truth", "Saved operator values are authoritative. AI suggestions are visibly distinct and do not silently become publish-ready inventory.", "success")
    doc.add_heading("Launch on Linux Mint", level=2)
    code = doc.add_paragraph()
    rr = code.add_run("cd ~/Projects/SnapIMS-v0.6.1\nsource .venv/bin/activate\nsnapims --data-dir ~/SnapIMS-data serve")
    rr.font.name = "Liberation Mono"
    rr.font.size = Pt(9)
    doc.add_paragraph("Open http://127.0.0.1:8767. Press Ctrl+C in the terminal to stop a foreground server.")
    doc.add_page_break()

    doc.add_heading("Physical capture and QR rules", level=1)
    doc.add_paragraph("CVHS1:ITEM:NEXT is the sole normal item boundary. Time gaps never split items. CVHS1:ITEM:CONT remains compatibility-only and performs no normal grouping action.")
    add_steps(doc, [
        "Photograph START first.",
        "Photograph the active shelf card A1-J10 or Q1.",
        "Photograph the VHS front first; it becomes the lead product image.",
        "Take any additional back, spine, tape, open-case, label, or damage photographs.",
        "Photograph NEXT after finishing the item and before the next tape.",
        "Photograph RARE and/or REVIEW after NEXT and before the flagged item front.",
        "Change shelf only between items.",
        "Photograph END last. A trailing NEXT is unnecessary.",
    ])
    qr = ASSETS / "qr-core.png"
    if qr.exists():
        add_picture(doc, qr, "Core command cards from Canada VHS QR Card Set v2", width=6.5)
    add_callout(doc, "Q1 quarantine", "Use Q1 for mold, damage, testing, or uncertainty. Do not clear a quarantined tape for listing until the physical issue is resolved.", "warning")
    doc.add_page_break()

    new_section(doc, "1. Home dashboard", "Home provides truthful workload counts and direct next actions without exposing Python internals.", [
        "Choose Import new batch for a new camera roll.",
        "Resume the current batch from Review or Batch Editor.",
        "Use Publish only after blockers are understood.",
        "Open Diagnostics when schema, provider, import, CSV, checkpoint, or transaction health is uncertain.",
    ], "01-home-production.png", "Home dashboard in production mode")

    new_section(doc, "2. Import preview", "Preview parses the chronological QR stream without creating permanent batch or item records.", [
        "Choose the configured incoming folder, a recent folder, or use the available folder-selection control.",
        "Enter an optional short batch label.",
        "Run Preview and compare item, product-photo, command, and warning totals with the physical session.",
        "Do not preserve a batch with unexplained warnings, wrong boundaries, or incorrect shelf assignment.",
    ], "02-import-preview.png", "Preview - not imported yet")

    new_section(doc, "3. Preserve and import", "Preserve and Import creates one durable Batch ID, immutable Item IDs, preserved originals, derivatives, manifests, and database records.", [
        "Click Preserve and Import once after Preview is correct.",
        "Wait for the durable Batch ID and imported totals.",
        "Continue to Review from the imported result.",
        "If the process is interrupted, restart SnapIMS and use Diagnostics; the import journal reconciles resumable, complete, failed, or quarantined work.",
    ], "03-import-complete.png", "Imported batch with durable identity")

    new_section(doc, "4. Provider safety", "Production mode shows only production-capable recognition providers. Mock, fixture, demo, and synthetic providers are disabled unless explicit test mode is enabled.", [
        "Confirm the intended provider is available in Settings or Review.",
        "Never treat a hidden or forged test-provider request as production recognition; the server rejects it.",
        "Historical test-sourced results remain visible for audit but cannot satisfy production publish readiness.",
        "Do not expose provider keys in screenshots, logs, tickets, or documentation.",
    ], "04-production-provider-selector.png", "Production provider selector - test providers absent")

    new_section(doc, "5. Recognition blocked and manual recovery", "Missing credentials or provider unavailability is a blocked pre-attempt state, not a fictional failed image recognition.", [
        "Read the visible blocker and remaining count.",
        "Configure a live provider and retry, or choose manual review.",
        "Manual review keeps the same Item ID and photographs and requires only the minimum operator fields.",
        "A manually completed item is marked Manual / No AI confidence rather than given invented confidence.",
    ], "05-recognition-blocked-manual-recovery.png", "Blocked recognition with clear manual recovery")

    new_section(doc, "6. Suggested, saved, and reviewed values", "SnapIMS distinguishes AI suggestions from durable working values and reviewed inventory.", [
        "AI suggestion means provider output that has not yet been accepted.",
        "Saved means the current durable operator value.",
        "Reviewed means the operator or an approved external-review workflow accepted the record.",
        "Drafted or Published refers to remote Shopify state and is never implied by a local save or rollback.",
        "Confidence shown on an unfinished item is latest suggestion confidence, not reviewed confidence.",
    ], "06-test-recognition-suggestion-state.png", "Suggestion state visibly separated from saved operator state")

    new_section(doc, "7. Fast Review: Price, Enter, next tape", "Routine Review is optimized for the keyboard. The Price field is selected on an unfinished item, so an experienced operator can type Price, press Enter, and continue.", [
        "Compare the current photograph with the displayed suggested or saved title.",
        "Optionally replace Price and/or Discount.",
        "Press Enter from the selected Price field or click Approve & Next.",
        "The current item is validated and saved atomically, unfinished count decreases, and the next unfinished tape opens exactly once.",
        "Rapid double Enter is guarded against approving two tapes.",
        "Use Later only for a genuine exception; it leaves the item unfinished.",
    ], "07-price-enter-advances-once.png", "Price then Enter advances exactly once")

    doc.add_heading("8. Keyboard reference", level=1)
    doc.add_paragraph("Visible controls remain authoritative; shortcuts reduce operator travel without hiding the workflow.")
    add_keyboard_table(doc)
    add_callout(doc, "Focus rule", "After Approve & Next, focus returns to the selected Price field on the next unfinished tape whenever the browser can do so safely.", "info")
    doc.add_page_break()

    new_section(doc, "9. Batch Editor and atomic bulk changes", "Batch Editor is for high-volume correction and exception triage. Suggested and saved values remain distinct, and all-or-nothing bulk actions never report false success.", [
        "Filter by confidence, status, flags, text, or location.",
        "Select the exact visible rows that should change.",
        "Open a bulk action such as Set Price, Set Discount, add/subtract, round, tags, flags, or location.",
        "Review the preview and affected count before applying.",
        "A successful action creates a checkpoint, writes audit history, and commits the selected operation atomically.",
        "If one row fails validation or persistence, the all-or-nothing action rolls back and reports exact failures.",
    ], "08-batch-editor-atomic-bulk.png", "Batch Editor after an atomic 20-row bulk operation")

    new_section(doc, "10. Command palette", "Ctrl+Shift+P opens the lightweight command palette. It supplements visible controls; it does not replace them.", [
        "Type to filter the available commands.",
        "Use it to open Review, Batch Editor, Publish, Diagnostics, CSV tools, confidence filters, or item jumps where advertised.",
        "Unavailable commands are not presented as functioning actions.",
        "Press Escape to close the palette.",
    ], "09-command-palette.png", "Command palette in Batch Editor")

    new_section(doc, "11. CSV download, stage, diff, and apply", "CSV is a deliberate spreadsheet editing surface keyed by immutable Item ID. Upload never silently overwrites inventory.", [
        "Download the current inventory_work.csv for the selected batch.",
        "Edit values without changing Item ID or moving data between batches.",
        "Upload the edited CSV. SnapIMS validates the entire file and stages it without changing authoritative records.",
        "Inspect the difference preview, validation errors, affected rows, and exact money values.",
        "Apply once. SnapIMS creates a rollback checkpoint, updates all rows and field history in one transaction, and marks the stage applied.",
        "Cancel changes nothing. A failed all-or-nothing apply leaves zero partial authoritative changes.",
    ], "10-csv-difference-preview.png", "Staged CSV difference preview before atomic apply")

    new_section(doc, "12. Shopify simulation and provenance blockers", "Publish uses saved working values only. Simulation is the default safe action, and production mode blocks test-sourced recognition from publish eligibility.", [
        "Choose the reviewed batch and inspect Ready, Blocked, Drafted, and Failed counts.",
        "Resolve any item where an AI suggestion exists but has not been accepted.",
        "Resolve any historical test-provider provenance by deliberate manual review or a live-provider result; do not silently relabel it.",
        "Run simulation before any live draft creation.",
        "For the eventual live acceptance test, create exactly one Shopify draft and verify Item ID/SKU, title, price, quantity, images, and admin link.",
        "Local checkpoint restore does not claim to undo a remote Shopify draft or publication.",
    ], "13-shopify-simulation-production-manual-replacement.png", "Production simulation after deliberate manual replacement of test provenance")

    new_section(doc, "13. Restart durability", "Saved inventory, recognition state, checkpoints, CSV stages, and incomplete import journals are durable across a full application-process restart.", [
        "Stop and restart the SnapIMS process.",
        "Open Home, Review, or Batch Editor and select the same batch.",
        "Verify Item IDs, titles, prices, discounts, shelves, flags, images, review state, and working-batch revision.",
        "Interrupted identifying work returns as a recoverable paused state; blocked work remains blocked; committed work is not duplicated.",
        "Stop if any durable value differs from the pre-restart record.",
    ], "12-restart-durability-production.png", "Same batch and Item IDs after a full application restart")

    new_section(doc, "14. Diagnostics and recovery", "Diagnostics is the operational truth source when a normal workflow cannot continue.", [
        "Confirm application version 0.6.1 and SQLite schema version 7.",
        "Require SQLite integrity check OK and zero foreign-key violations.",
        "Review schema-manifest status, database/WAL size, last backup, checkpoints, staged CSV files, import journals, and current revision.",
        "Inspect provider configuration state without exposing secrets.",
        "Use the documented recovery action; do not bypass schema failure, change IDs, or delete the operator database to make the screen green.",
    ], "14-diagnostics.png", "Expanded system integrity and recovery diagnostics")

    doc.add_heading("End-of-batch checklist", level=1)
    add_bullets(doc, [
        "Imported item count matches the physical VHS count.",
        "Every item has the correct lead image and additional photographs.",
        "Item ID and Batch ID remained unchanged through all corrections.",
        "Shelf/location and RARE / Physical Review flags match the physical tape.",
        "All Later items were revisited or deliberately remain unfinished.",
        "Every saved title, Price, Discount, quantity, and condition is intentional.",
        "Suggested and saved values are not confused.",
        "Low-confidence and failed-recognition exceptions are resolved or clearly deferred.",
        "CSV diff/apply and rollback states are understood.",
        "Publish simulation shows expected Ready and Blocked counts.",
        "Test-sourced records are not production-publish eligible.",
        "Original camera photographs remain untouched.",
        "Diagnostics reports schema 7, integrity OK, and zero foreign-key violations.",
    ])
    add_callout(doc, "1.0 remains forbidden", "Do not label SnapIMS 1.0.0 until the real 20-tape Pixel pilot, live AI, one live Shopify draft, physical CSV verification, restart durability, Operator Guide walkthrough, browser verification, and blocker review all pass.", "danger")
    doc.add_heading("External boundaries", level=2)
    add_bullets(doc, [
        "Browser fixtures verify workflow behaviour but do not prove live AI quality, latency, billing, or provider limits.",
        "Shopify simulation does not prove live authentication, staged media upload, inventory mutation, or draft creation.",
        "The current architecture is a single-operator local workstation. Remote exposure does not make it multi-user safe.",
        "Warehouse-scale 5,000-row virtualization, adaptive recognition, near-duplicate classification, and complete remote-operation security remain deferred.",
    ])

    core = doc.core_properties
    core.title = f"SnapIMS Operator Guide v{VERSION}"
    core.subject = "Photo-first VHS inventory operator procedures"
    core.author = "Canada VHS"
    core.keywords = "SnapIMS, Canada VHS, operator guide, inventory, QR, Shopify"
    core.comments = "Generated from the final browser-verified v0.6.1 UI."
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
