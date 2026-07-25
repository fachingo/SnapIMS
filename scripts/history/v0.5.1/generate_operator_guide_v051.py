from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
SHOTS = ROOT / "operator-audit-assets" / "v0.5.1-browser" / "screenshots"
OUT = ROOT / "SnapIMS_Operator_Guide.docx"

PAGES = [
    ("Home", "Start here. Import a new camera roll or continue the active batch. The summary cards show current workload.", "00-home.png"),
    ("Import - first run", "Use Advanced once to choose the camera-roll folder. Save it as the incoming folder so routine later imports use the selector.", "01-first-run-import.png"),
    ("Preview - not imported yet", "Confirm item, product-photo, command, and warning counts. Preview is temporary and does not display a durable Batch ID.", "02-preview-not-imported.png"),
    ("Preserve and import", "After grouping is correct, preserve/import. SnapIMS displays exactly one durable Batch ID and leaves source photographs untouched.", "02-imported-durable-id.png"),
    ("Review before identification", "The compact status explains the next action. Select a provider and choose Identify items once.", "04-review-before-identification.png"),
    ("Recognition complete", "The suggested title, Price, Discount, physical position, and primary action are visible together.", "05-recognition-complete.png"),
    ("Quick edit - Price", "Enter a corrected CAD price directly in Review, then choose Approve & Next. No separate Save is required.", "06-price-quick-edit.png"),
    ("Quick edit - Discount", "Enter the percentage directly in Review, then choose Approve & Next once.", "07-discount-quick-edit.png"),
    ("Quick edit - Price and Discount", "Both fields can be changed before the same one-action approval.", "08-price-discount-quick-edit.png"),
    ("Automatic advancement", "Approve & Next marks the current item Done, decrements unfinished count, and opens the next physical item exactly once.", "09-next-item-opened.png"),
    ("Review complete", "When unfinished count reaches zero, the Review-complete message replaces normal recognition actions. Use Done to inspect records.", "10-review-complete.png"),
    ("Correct a completed item", "Edit item opens saved data. Save changes stays on the same immutable Item ID.", "11-completed-item-corrected.png"),
    ("Validation blocks invalid changes", "Blank Title and zero Price are rejected atomically. The editor remains open with exact errors.", "12-invalid-edit-blocked.png"),
    ("Cancel invalid changes", "Cancel changes restores the last valid saved record without altering the Item ID.", "13-invalid-edit-cancelled.png"),
    ("Restart durability", "After a full application-process restart, saved metadata and Item ID remain available from SQLite.", "14-restart-durable-correction.png"),
    ("Publish simulation", "Confirm Ready/Blocked counts and payload previews. Download the inventory CSV before any live-draft acceptance test.", "15-publish-simulation.png"),
    ("Routine second import - Preview", "Configured and recent folders avoid routine absolute-path typing.", "16-preview-not-imported.png"),
    ("Routine second import - durable ID", "A successful import is available immediately through Continue to Review.", "16-imported-durable-id.png"),
    ("Postpone with Later", "Later does not finish the tape. The notice confirms it remains unfinished; return through Unresolved.", "18-later-preserves-unfinished.png"),
    ("Missing incoming folder", "SnapIMS fails closed and does not scan another folder. Reconnect the device or use Advanced recovery.", "19-missing-folder-recovery.png"),
    ("Recognition failure fixture - Preview", "The normal Preview contract remains unchanged for a one-item failure test.", "20-preview-not-imported.png"),
    ("Recognition failure fixture - Import", "The failed provider attempt remains attached to this same durable item.", "20-imported-durable-id.png"),
    ("Recognition failure", "Open the Failed queue, read the provider error, choose a working provider, and Retry.", "22-recognition-failure.png"),
    ("Recognition recovered", "Retry restores a reviewable suggestion without creating another item or changing physical position.", "23-recognition-failure-recovered.png"),
    ("Interruption fixture - Preview", "A large deterministic batch keeps recognition visibly active long enough to test process interruption.", "24-preview-not-imported.png"),
    ("Interruption fixture - Import", "The 40-item batch receives one durable identity before recognition starts.", "24-imported-durable-id.png"),
    ("Recognition running", "The browser shows committed progress before the application process is terminated.", "26-recognition-running-before-kill.png"),
    ("Recognition paused after restart", "Startup converts the orphaned Running job to Paused with exact completed/remaining counts and one Continue action.", "27-recognition-paused-after-restart.png"),
    ("Continue interrupted identification", "Continue resumes from committed boundaries and completes without duplicate recognition attempts.", "28-recognition-resumed-complete.png"),
    ("Settings", "Set the normal incoming folder, inspect recent folders, and confirm provider availability before a real pilot.", "29-settings.png"),
    ("Diagnostics", "SQLite integrity must be ok, foreign-key violations zero, and schema version 5 before operation or restore.", "30-diagnostics.png"),
]


def field(run, instruction: str) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, sep, end])


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_page(doc: Document, title: str, body: str, image_name: str) -> None:
    doc.add_page_break()
    p = doc.add_paragraph()
    p.style = doc.styles["Title"]
    p.add_run(title)
    p = doc.add_paragraph(body)
    p.style = doc.styles["Normal"]
    p.paragraph_format.space_after = Pt(8)
    image_path = SHOTS / image_name
    with Image.open(image_path) as image:
        ratio = image.width / image.height
    max_w, max_h = 6.55, 6.35
    if ratio >= max_w / max_h:
        width = max_w
        height = width / ratio
    else:
        height = max_h
        width = height * ratio
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(image_path), width=Inches(width), height=Inches(height))
    caption = doc.add_paragraph(image_name)
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.style = doc.styles["Caption"]


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.55)
section.bottom_margin = Inches(0.55)
section.left_margin = Inches(0.7)
section.right_margin = Inches(0.7)

styles = doc.styles
styles["Normal"].font.name = "Liberation Sans"
styles["Normal"].font.size = Pt(10.5)
styles["Title"].font.name = "Liberation Sans"
styles["Title"].font.size = Pt(24)
styles["Title"].font.bold = True
styles["Title"].font.color.rgb = RGBColor(27, 119, 79)
styles["Heading 1"].font.name = "Liberation Sans"
styles["Heading 1"].font.color.rgb = RGBColor(27, 119, 79)
styles["Heading 2"].font.name = "Liberation Sans"
styles["Heading 2"].font.color.rgb = RGBColor(27, 119, 79)
styles["Caption"].font.name = "Liberation Sans"
styles["Caption"].font.size = Pt(8)

# Footer and header.
header = section.header.paragraphs[0]
header.text = "SnapIMS 0.5.1 Operator Guide"
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
header.runs[0].font.size = Pt(8)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer.add_run("SnapIMS 0.5.1 Operator Guide  |  Page ")
r.font.size = Pt(8)
field(footer.add_run(), "PAGE")

# Cover.
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Inches(1.5)
r = p.add_run("SnapIMS")
r.bold = True
r.font.size = Pt(40)
r.font.color.rgb = RGBColor(27, 119, 79)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Operator Guide")
r.bold = True
r.font.size = Pt(28)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Version 0.5.1 - Verification Hardening Patch").font.size = Pt(16)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Photo-first inventory ingestion for Canada VHS\nFastAPI/Jinja browser workstation · SQLite schema 5")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Inches(1.2)
p.add_run("Use the final browser UI as the source of truth.\nMock and simulation screenshots do not prove live AI or live Shopify.")

# Quick start.
doc.add_page_break()
doc.add_heading("Quick start", level=1)
doc.add_paragraph("The routine operator flow is intentionally short.")
table = doc.add_table(rows=2, cols=5)
table.style = "Table Grid"
headers = ["1. Photograph", "2. Import", "3. Identify", "4. Review", "5. Publish"]
values = [
    "START, shelf, tape photos, NEXT, repeat, END.",
    "Preview grouping, then preserve/import.",
    "Choose a provider and identify the batch.",
    "Confirm title; optionally adjust Price or Discount; Approve & Next.",
    "Simulate drafts and download the CSV.",
]
for i, text in enumerate(headers):
    table.rows[0].cells[i].text = text
    set_cell_shading(table.rows[0].cells[i], "D9EDE3")
for i, text in enumerate(values):
    table.rows[1].cells[i].text = text
for bullet in [
    "Edit is exception mode, not the routine path.",
    "Later postpones a tape and leaves it unfinished.",
    "Never change Item ID; it is permanent across Review, CSV, and Shopify.",
    "Real operator time per tape must be measured during the live pilot.",
]:
    doc.add_paragraph(bullet, style="List Bullet")

doc.add_heading("QR photography sequence", level=2)
qr = doc.add_table(rows=1, cols=3)
qr.style = "Table Grid"
for i, text in enumerate(["Card / action", "Payload", "Meaning"]):
    qr.rows[0].cells[i].text = text
    set_cell_shading(qr.rows[0].cells[i], "D9EDE3")
rows = [
    ("START", "CVHS1:BATCH:START", "Begins the batch."),
    ("LOCATION", "CVHS1:LOC:A1 ... J10 or Q1", "Applies shelf to following items."),
    ("PRODUCT", "ordinary photos", "First product photo is Front."),
    ("NEXT", "CVHS1:ITEM:NEXT", "Sole normal item boundary."),
    ("FLAGS", "CVHS1:FLAG:RARE / REVIEW", "Apply to next item, then reset."),
    ("END", "CVHS1:BATCH:END", "Closes final item and batch."),
]
for a, b, c in rows:
    cells = qr.add_row().cells
    cells[0].text, cells[1].text, cells[2].text = a, b, c

doc.add_paragraph("Time gaps never split items. CONT is deprecated compatibility only. Command images are preserved for audit but never attached to products.")

# One page per screenshot.
for title, body, image in PAGES:
    add_page(doc, title, body, image)

# Final checklist.
doc.add_page_break()
doc.add_heading("End-of-batch checklist", level=1)
checks = [
    "Unfinished count is zero and Review shows Review complete.",
    "Every title matches the photograph.",
    "Quick-edited Price and Discount values are intentional.",
    "Shelf and Rare / Physical Review flags match the physical tape.",
    "Completed corrections retain the original Item ID.",
    "Publish simulation shows expected Ready and Blocked counts.",
    "The browser-downloaded CSV contains the same Item IDs and saved titles.",
    "During the first real pilot, keep Shopify in simulation until physical CSV reconciliation passes.",
    "For the live-draft acceptance test, create exactly one draft and inspect it in Shopify.",
]
for item in checks:
    doc.add_paragraph(f"☐ {item}")

doc.add_heading("1.0 acceptance gate", level=2)
doc.add_paragraph(
    "Version 1.0.0 is forbidden until the real 20-tape Pixel pilot, live AI, one live Shopify draft, physical CSV verification, restart durability, independent guide walkthrough, browser verification, and blocker review all pass."
)

doc.core_properties.title = "SnapIMS 0.5.1 Operator Guide"
doc.core_properties.subject = "Canada VHS photo-first inventory workflow"
doc.core_properties.author = "SnapIMS Project"
doc.core_properties.comments = "Generated from the verified v0.5.1 browser UI."
doc.save(OUT)
print(OUT)
