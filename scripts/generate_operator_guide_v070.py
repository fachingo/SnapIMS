from __future__ import annotations

import math
from pathlib import Path
from typing import Iterable

from PIL import Image
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SHOTS = ROOT / "release-evidence" / "v0.7.0" / "browser" / "screenshots"
OUT = ROOT / "SnapIMS_Operator_Guide_v0.7.0.docx"
ASSETS = ROOT / ".guide-assets-v070"
VERSION = "0.7.0"

ACCENT = "8B5CF6"
DARK = "17131F"
MUTED = "676071"
LIGHT = "F3EEFC"
GREEN = "176B4D"
YELLOW = "8A5A00"
RED = "A9364D"
BORDER = "D9D2E6"


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def border(cell, color: str = BORDER) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), color)


def page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(f"SnapIMS Operator Guide v{VERSION}  |  ")
    run.font.name = "Liberation Sans"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t"); text.text = "1"
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, sep, text, end])


def configure(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.48)
    section.bottom_margin = Inches(0.48)
    section.left_margin = Inches(0.62)
    section.right_margin = Inches(0.62)
    section.header_distance = Inches(0.15)
    section.footer_distance = Inches(0.18)
    page_number(section.footer.paragraphs[0])

    normal = doc.styles["Normal"]
    normal.font.name = "Liberation Sans"
    normal.font.size = Pt(9.3)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.02
    for name, size, color in [
        ("Title", 29, DARK), ("Subtitle", 12, MUTED),
        ("Heading 1", 18, ACCENT), ("Heading 2", 12.5, DARK), ("Heading 3", 10.5, ACCENT),
    ]:
        style = doc.styles[name]
        style.font.name = "Liberation Sans"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = name != "Subtitle"
        style.paragraph_format.space_after = Pt(5)


def callout(doc: Document, title: str, body: str, kind: str = "info") -> None:
    fills = {"info": LIGHT, "success": "E9F6F0", "warning": "FFF4D7", "danger": "FCE8EC"}
    accents = {"info": ACCENT, "success": GREEN, "warning": YELLOW, "danger": RED}
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    shade(cell, fills[kind]); border(cell)
    p = cell.paragraphs[0]
    r = p.add_run(title); r.bold = True; r.font.color.rgb = RGBColor.from_string(accents[kind])
    p2 = cell.add_paragraph(body); p2.paragraph_format.space_after = Pt(0)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def steps(doc: Document, values: Iterable[str]) -> None:
    for i, value in enumerate(values, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.first_line_indent = Inches(-0.2)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(f"{i}. "); r.bold = True; r.font.color.rgb = RGBColor.from_string(ACCENT)
        p.add_run(value)


def bullets(doc: Document, values: Iterable[str]) -> None:
    for value in values:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2.5)
        p.add_run(value)


def picture(doc: Document, filename: str, caption: str, width: float = 7.0) -> None:
    path = SHOTS / filename
    if not path.exists():
        raise FileNotFoundError(path)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Inches(width))
    cp = doc.add_paragraph(caption); cp.alignment = WD_ALIGN_PARAGRAPH.CENTER; cp.paragraph_format.space_after = Pt(3)
    for run in cp.runs:
        run.italic = True; run.font.size = Pt(7.5); run.font.color.rgb = RGBColor.from_string(MUTED)



def screenshot_slices(filename: str, max_height: int = 1080) -> list[Path]:
    path = SHOTS / filename
    if not path.exists():
        raise FileNotFoundError(path)
    ASSETS.mkdir(exist_ok=True)
    image = Image.open(path).convert("RGB")
    if image.height <= max_height:
        out = ASSETS / f"guide-{filename}"
        image.save(out, optimize=True)
        return [out]
    overlap = 70
    part_count = math.ceil(image.height / max_height)
    base = math.ceil(image.height / part_count)
    parts: list[Path] = []
    for index in range(part_count):
        top = max(0, index * base - (overlap if index else 0))
        bottom = min(image.height, (index + 1) * base + (overlap if index < part_count - 1 else 0))
        crop = image.crop((0, top, image.width, bottom))
        out = ASSETS / f"guide-{Path(filename).stem}-part-{index + 1:02d}.png"
        crop.save(out, optimize=True)
        parts.append(out)
    return parts


def screenshot_pages(doc: Document, title: str, filename: str, caption: str) -> None:
    parts = screenshot_slices(filename)
    for index, part in enumerate(parts, 1):
        doc.add_heading(title if len(parts) == 1 else f"{title} — screen {index} of {len(parts)}", level=2)
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(2)
        p.add_run().add_picture(str(part), width=Inches(7.0))
        cp = doc.add_paragraph(caption if len(parts) == 1 else f"{caption} — view {index} of {len(parts)}")
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER; cp.paragraph_format.space_after = Pt(2)
        for run in cp.runs:
            run.italic = True; run.font.size = Pt(7.5); run.font.color.rgb = RGBColor.from_string(MUTED)
        doc.add_page_break()

def screen_page(doc: Document, title: str, intro: str, items: list[str], image: str, caption: str, note: tuple[str, str, str] | None = None) -> None:
    doc.add_heading(title, level=1)
    doc.add_paragraph(intro)
    steps(doc, items)
    if note:
        callout(doc, note[0], note[1], note[2])
    doc.add_page_break()
    screenshot_pages(doc, title, image, caption)


def shortcut_table(doc: Document) -> None:
    rows = [
        ("Ctrl/Cmd+Shift+P", "Open Quick Command Palette; Ctrl/Cmd+K is the browser-safe fallback."),
        ("Enter in Review", "Submit valid Title/Price/Discount once and open the next unfinished tape."),
        ("Arrow Up/Down", "Move to the same editable Batch Editor field in the previous/next visible row."),
        ("Arrow Left/Right", "Move between editable cells when the text caret is already at a boundary."),
        ("Enter in Batch Editor", "Save the current cell and move down in the same field."),
        ("Shift+Arrow", "Extend visible-row selection."),
        ("Space", "Toggle the active row when not typing in an editor."),
        ("Ctrl/Cmd+A", "Select all currently visible filtered rows while the Batch Editor has focus."),
        ("Ctrl/Cmd+1–9", "Open the first nine quick actions through the normal preview dialog."),
        ("Alt+Left/Right", "Reorder the focused quick action."),
        ("Escape", "Cancel an edit/dialog, close the palette, or clear selection as context allows."),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.columns[0].width = Inches(1.65); table.columns[1].width = Inches(5.25)
    for idx, text in enumerate(("Shortcut", "Action")):
        cell = table.rows[0].cells[idx]; cell.text = text; shade(cell, ACCENT); border(cell, ACCENT)
        for run in cell.paragraphs[0].runs:
            run.bold = True; run.font.color.rgb = RGBColor(255,255,255)
    for key, action in rows:
        cells = table.add_row().cells; cells[0].text = key; cells[1].text = action
        for cell in cells:
            border(cell); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cells[0].paragraphs[0].runs[0].bold = True


def build() -> None:
    doc = Document(); configure(doc)

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(32)
    r = p.add_run("SnapIMS"); r.font.name = "Liberation Sans"; r.font.size = Pt(42); r.font.bold = True; r.font.color.rgb = RGBColor.from_string(ACCENT)
    p = doc.add_paragraph("Operator Guide"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.style = doc.styles["Title"]
    p = doc.add_paragraph("Version 0.7.0 — Keyboard Workstation and Local Movie Catalog"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.style = doc.styles["Subtitle"]
    callout(doc, "Operating model", "SnapIMS is an exception-handling workstation: inspect the photograph, confirm or type Title, optionally change Price/Discount, then press Enter once.", "info")
    callout(doc, "Release boundary", "Version 0.7.0 is not 1.0.0. Live AI, one live Shopify draft, physical CSV verification, the real Pixel pilot, production restart, guide walkthrough and blocker review remain mandatory.", "warning")
    picture(doc, "01-home-production.png", "SnapIMS 0.7.0 Home — final native-browser evidence", 6.8)
    doc.add_page_break()

    doc.add_heading("Quick start, installation and safety", level=1)
    steps(doc, [
        "Photograph START, shelf, the tape front and support photos, NEXT between tapes, and END last.",
        "Preview the complete camera folder and preserve only when counts, boundaries and warnings make sense.",
        "Identify with a configured live provider or continue manually.",
        "In Review, type Title only when empty; otherwise inspect the suggested/saved Title. Adjust Price/Discount if needed and press Enter once.",
        "Use Batch Editor, checkpoints or staged CSV for deliberate batch work.",
        "Run Shopify simulation before any live draft acceptance test.",
    ])
    doc.add_heading("Launch on Linux Mint", level=2)
    code = doc.add_paragraph(); run = code.add_run("cd ~/Projects/SnapIMS-v0.7.0\nsource .venv/bin/activate\nsnapims --data-dir ~/SnapIMS-data serve")
    run.font.name = "Liberation Mono"; run.font.size = Pt(9)
    doc.add_paragraph("Open http://127.0.0.1:8767. Install the optional native folder picker with sudo apt install python3-tk. Without it, SnapIMS exposes the manual-path fallback.")
    callout(doc, "Identity rule", "Never relink by title, filename, spreadsheet row, provider ID or sort order. Item ID and Batch ID are permanent.", "danger")
    doc.add_page_break()

    doc.add_heading("Physical capture and QR rules", level=1)
    doc.add_paragraph("CVHS1:ITEM:NEXT is the sole normal item boundary. Time gaps never split items. ITEM:CONT remains compatibility-only.")
    steps(doc, [
        "Photograph CVHS1:BATCH:START first.",
        "Photograph the active shelf A1–J10 or quarantine Q1.",
        "Photograph the VHS front first; it becomes the lead image.",
        "Take back, spine, tape, label, damage or open-case photos as needed.",
        "Photograph NEXT after the item and before the next front.",
        "After NEXT, photograph RARE and/or REVIEW flags for the upcoming tape.",
        "Change shelf only between items.",
        "Photograph END last. A trailing NEXT is not required.",
    ])
    callout(doc, "Q1 quarantine", "Use Q1 for mold, damage, testing or uncertainty. Do not clear the tape for listing until the physical issue is resolved.", "warning")
    callout(doc, "Originals", "Import preserves original camera images. Do not rename, edit or delete them as a normal workflow shortcut.", "success")
    doc.add_page_break()

    screen_page(doc, "1. Import preview", "Preview parses the QR stream without creating durable inventory.", [
        "Choose the configured/recent folder, use Browse Folder, or open Advanced and paste a manual path.",
        "If Tkinter is unavailable, read the brief fallback notice and paste the folder path; the workflow remains usable.",
        "Compare item, product-photo, command and warning totals with the physical session.",
        "Do not preserve unexplained warnings or incorrect item boundaries.",
    ], "02-import-preview.png", "20-item Preview — not imported yet")

    screen_page(doc, "2. Preserve and import", "Preserve creates one durable Batch ID, immutable Item IDs, preserved media and an import journal.", [
        "Click Preserve and Import once after Preview is correct.",
        "Record the durable Batch ID and confirm imported totals.",
        "Continue directly to Review.",
        "After interruption, restart and use Diagnostics; incomplete import states reconcile or quarantine deterministically.",
    ], "03-import-complete.png", "Durable Import result")

    screen_page(doc, "3. Review with an empty Title", "An empty Title is a normal quick-field exception, not a reason for extra Enter presses.", [
        "SnapIMS focuses Title because it is the first required empty quick field.",
        "Type the sale Title and optionally adjust Price or Discount.",
        "Press Enter once. The server validates and saves the entire quick record, marks the item Done and opens the next unfinished tape.",
        "The full exception editor appears only if genuine validation/conflict work remains.",
    ], "04-empty-title-one-enter.png", "Manual Title completed through the one-Enter path", ("Routine rule", "Title → Enter is Approve & Next when the quick record is valid.", "success"))

    screen_page(doc, "4. Movie ambiguity is an exception", "The Local Movie Catalog never silently chooses between materially ambiguous titles, years or media types.", [
        "Review candidate title, year, media type, director/source evidence and match score.",
        "Choose Use this movie only when the candidate is correct.",
        "Retry or leave the item unlinked when evidence is insufficient.",
        "Movie selection changes no physical Item ID, Batch ID, photo, shelf or Shopify identity.",
    ], "05-catalog-ambiguity-preserved.png", "Ambiguous candidate held for operator selection")

    screen_page(doc, "5. Local Movie match", "After operator confirmation, the Movie is stored locally with source provenance and reused for future matching.", [
        "Confirm the compact Catalog match label in Review.",
        "Use detailed catalog controls only for ambiguity, mismatch or recovery.",
        "The background job must never delay the next Review tape.",
        "Wikipedia supplies structured movie facts, not VHS packaging/edition authority or storefront artwork.",
    ], "05b-local-movie-catalog-match.png", "Operator-selected Movie persisted in the local catalog", ("Local first", "Repeated copies search the local catalog before any external request.", "info"))

    screen_page(doc, "6. Recognition blocked and manual recovery", "A missing key/provider is BLOCKED before recognition, not a fabricated failed identification.", [
        "Configure a live provider and retry, or continue with manual Review.",
        "Production mode does not display Mock beside live providers.",
        "Historical test-sourced data remains auditable and cannot silently satisfy live publish readiness.",
        "Never expose provider secrets in screenshots, logs or support reports.",
    ], "06-recognition-blocked-manual-recovery.png", "Blocked provider state with a real manual path")

    screen_page(doc, "7. Price → Enter → next tape", "When Title exists, Price is automatically focused and selected for the most common operator correction.", [
        "Inspect the photograph and displayed Title.",
        "Type a new Price, or leave the current saved value.",
        "Optionally adjust Discount.",
        "Press Enter once; the item saves and the next unfinished tape opens exactly once.",
        "Rapid double Enter is guarded against approving two records.",
    ], "07-price-enter-advances-once.png", "Price-selected one-Enter automatic advancement")

    screen_page(doc, "8. Quick Command Palette", "The palette supplements visible controls and opens from normal pages.", [
        "Press Ctrl/Cmd+Shift+P. Use Ctrl/Cmd+K if the browser reserves the primary combination.",
        "Type to filter available commands.",
        "Execute only commands that are enabled in the current context.",
        "Press Escape to close; search receives focus on open.",
    ], "08-command-palette.png", "Quick Command Palette")

    doc.add_heading("9. Batch Editor keyboard reference", level=1)
    doc.add_paragraph("The Batch Editor remains a smooth autosaving grid. Keyboard movement respects visible rows, filters, sorting, read-only cells and text-caret boundaries.")
    shortcut_table(doc)
    callout(doc, "Selection scope", "Ctrl/Cmd+A means currently visible filtered rows. Bulk actions must never silently include rows the operator could not reasonably know were selected.", "warning")
    doc.add_page_break()
    screenshot_pages(doc, "9. Batch Editor keyboard reference", "09-batch-editor-keyboard-grid.png", "Keyboard grid, visible-row selection and quick-action preview")

    screen_page(doc, "10. Quick actions and atomic bulk work", "The first nine visible actions have Ctrl/Cmd+1–9 shortcuts and use the same safe preview as a click.", [
        "Select the intended visible rows.",
        "Open an action by button or shortcut and read its concise effect/checkpoint/rollback description.",
        "Review the affected count and value before applying.",
        "Reorder actions by drag or Alt+Left/Right; Reset restores the default order.",
        "All-or-nothing operations roll back and report exact failures instead of claiming partial success.",
    ], "10-batch-editor-atomic-bulk.png", "Atomic 20-row bulk Price operation", ("Checkpoint", "Destructive or large bulk actions create a rollback checkpoint before commit.", "success"))

    screen_page(doc, "11. CSV stage, diff and apply", "CSV remains keyed by immutable Item ID and never overwrites inventory immediately on upload.", [
        "Download inventory_work.csv from the selected batch.",
        "Edit values without removing/changing Item ID or mixing batches.",
        "Upload the exact export or a spreadsheet-resaved comma/semicolon/tab file; harmless BOM/whitespace/header variation is normalized.",
        "Inspect detected columns, validation, difference preview and affected rows.",
        "Apply once to create a checkpoint and commit all changes atomically; use rollback when required.",
    ], "11-csv-difference-preview.png", "Staged CSV difference preview")

    screen_page(doc, "12. Shopify simulation", "Publish uses saved working values and structured local Movie facts; simulation remains the default safe action.", [
        "Inspect Ready, Blocked, Drafted and Failed counts.",
        "Resolve unsaved suggestions, test provenance and other blockers.",
        "Run simulation and inspect Item ID/SKU, Title, Price, quantity, images and structured Movie fields.",
        "Live mode creates drafts only and remains deliberately separate.",
        "A local rollback never claims to undo a remote Shopify draft/product.",
    ], "14-shopify-simulation-production-manual-replacement.png", "Production-safe Shopify simulation")

    screen_page(doc, "13. Restart durability", "Committed inventory, Review state, CSV stages, checkpoints, import journals and catalog jobs survive a full application-process restart.", [
        "Stop SnapIMS with Ctrl+C and restart the same data directory.",
        "Reopen the Batch from Home or the selectors.",
        "Verify Item IDs, titles, prices, discounts, shelf, flags, images, Review state and Movie link.",
        "Interrupted catalog/recognition jobs return as paused/retryable states without duplicate durable records.",
        "Stop if any committed value differs from the pre-restart record.",
    ], "13-restart-durability-production.png", "Same batch after full process restart")

    screen_page(doc, "14. Diagnostics and catalog administration", "Diagnostics is the operational truth source when a normal workflow cannot continue.", [
        "Require application 0.7.0, inventory schema 8, integrity OK, zero foreign-key violations and valid schema manifest.",
        "Inspect database/WAL size, backups, checkpoints, CSV stages, import journals and operation requests.",
        "Inspect Movie count, lookup jobs, candidates, failures, cache and FTS health.",
        "Use Verify, Backup, Restore, Retry, Rebuild FTS or Reconcile only with the documented recovery procedure.",
        "Never delete the operator database or change IDs just to make Diagnostics green.",
    ], "15-diagnostics-catalog.png", "Inventory and Local Movie Catalog diagnostics")

    doc.add_heading("End-of-batch checklist", level=1)
    bullets(doc, [
        "Imported item count matches the physical VHS count.",
        "Every Item has the correct lead image and support photos.",
        "Item ID and Batch ID remained unchanged.",
        "Shelf/location and RARE / Physical Review flags match the tape.",
        "All Later items were revisited or deliberately remain unfinished.",
        "Every saved Title, Price, Discount, quantity and condition is intentional.",
        "Ambiguous/mismatched Movie candidates are resolved or explicitly left unlinked.",
        "CSV stage/apply/rollback state is understood and reconciled.",
        "Shopify simulation shows the expected Ready/Blocked counts and no automatic publication occurred.",
        "Original camera images remain untouched.",
        "Diagnostics reports schema 8, inventory/catalog integrity OK and zero foreign-key violations.",
        "A catalog backup exists before catalog restore/rebuild work.",
    ])
    callout(doc, "1.0.0 remains forbidden", "Do not label SnapIMS 1.0.0 until the real 20-tape Pixel pilot, live AI, one live Shopify draft, physical CSV verification, production restart durability, independent guide walkthrough, browser verification and blocker review all pass.", "danger")
    doc.add_heading("Current external boundaries", level=2)
    bullets(doc, [
        "Browser fixtures verify workflow behaviour but do not prove live AI quality, billing or provider limits.",
        "The default test suite uses bounded Wikipedia fixtures; the live integration test is opt-in.",
        "Shopify simulation does not prove live authentication, media upload or draft creation.",
        "The current release is a single-operator workstation and does not claim 5,000-row/10,000-photo readiness.",
    ])

    props = doc.core_properties
    props.title = f"SnapIMS Operator Guide v{VERSION}"
    props.subject = "Canada VHS photo-first inventory operator procedures"
    props.author = "Canada VHS"
    props.keywords = "SnapIMS, Canada VHS, VHS, inventory, QR, local movie catalog, Wikipedia, Shopify"
    props.comments = "Generated from the final browser-verified SnapIMS 0.7.0 UI."
    doc.save(OUT)
    (ROOT / "SnapIMS_Operator_Guide.docx").write_bytes(OUT.read_bytes())
    print(OUT)


if __name__ == "__main__":
    build()
