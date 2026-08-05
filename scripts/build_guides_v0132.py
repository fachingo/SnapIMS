#!/usr/bin/env python3
from __future__ import annotations
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.shared import Inches, Pt
from PIL import Image
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
DOCS=ROOT/'docs'
EVIDENCE=ROOT/'browser-evidence'/'v0.13.2'
VERSION='0.13.2'
DATE='August 4, 2026'


def set_cell_shading(cell, fill='E8EEF7'):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)

def add_page_number(paragraph):
    paragraph.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    run=paragraph.add_run('Page ')
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); run._r.addnext(fld)

def setup(title, subtitle):
    d=Document()
    sec=d.sections[0]
    sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.75); sec.right_margin=Inches(.75)
    styles=d.styles
    styles['Normal'].font.name='Arial'; styles['Normal'].font.size=Pt(10)
    styles['Title'].font.name='Arial'; styles['Title'].font.size=Pt(24); styles['Title'].font.bold=True
    for name,size in [('Heading 1',17),('Heading 2',13),('Heading 3',11)]:
        styles[name].font.name='Arial'; styles[name].font.size=Pt(size); styles[name].font.bold=True
    p=d.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run(title)
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(subtitle); r.bold=True; r.font.size=Pt(13)
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run(f'Active release v{VERSION} · {DATE}')
    d.add_paragraph('SnapIMS is local-first inventory software. This guide describes the browser interface and command behavior verified for the active release.', style='Intense Quote')
    d.add_page_break()
    for section in d.sections:
        header=section.header.paragraphs[0]; header.text=f'SnapIMS v{VERSION} · {title}'; header.alignment=WD_ALIGN_PARAGRAPH.CENTER
        add_page_number(section.footer.paragraphs[0])
    return d

def h(d,text,level=1): d.add_heading(text,level=level)
def p(d,text,bold_prefix=None):
    para=d.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        para.add_run(bold_prefix).bold=True; para.add_run(text[len(bold_prefix):])
    else: para.add_run(text)
    return para

def bullets(d,items):
    for item in items: d.add_paragraph(item,style='List Bullet')

def numbered(d,items):
    for item in items: d.add_paragraph(item,style='List Number')

def code(d,text):
    para=d.add_paragraph(); para.style=d.styles['Normal'];
    run=para.add_run(text); run.font.name='Courier New'; run.font.size=Pt(9)
    para.paragraph_format.left_indent=Inches(.25); para.paragraph_format.space_before=Pt(3); para.paragraph_format.space_after=Pt(6)

def image(d,name,caption):
    path=EVIDENCE/name
    if path.exists():
        with Image.open(path) as im:
            width_px, height_px = im.size
        ratio = width_px / max(1, height_px)
        max_w, max_h = 6.6, 6.7
        width = min(max_w, max_h * ratio)
        height = width / ratio
        d.add_picture(str(path), width=Inches(width), height=Inches(height))
        q=d.paragraphs[-1]; q.alignment=WD_ALIGN_PARAGRAPH.CENTER
        c=d.add_paragraph(caption); c.alignment=WD_ALIGN_PARAGRAPH.CENTER; c.runs[0].italic=True

def table(d,headers,rows,widths=None):
    t=d.add_table(rows=1,cols=len(headers)); t.style='Table Grid'
    for i,x in enumerate(headers):
        t.rows[0].cells[i].text=str(x); set_cell_shading(t.rows[0].cells[i])
        for r in t.rows[0].cells[i].paragraphs[0].runs:r.bold=True
    for row in rows:
        cells=t.add_row().cells
        for i,x in enumerate(row): cells[i].text=str(x)
    return t

def save(d,name):
    DOCS.mkdir(parents=True,exist_ok=True); path=DOCS/name; d.save(path); print(path)


def day_to_day():
    d=setup('SnapIMS Day-to-Day Guide','First-time operator workflow')
    h(d,'1. Start SnapIMS')
    code(d,'export PATH="$HOME/.local/bin:$PATH"\nsnapims up\nsnapims status')
    bullets(d,[
        'Version must show 0.13.2.',
        'Inventory schema must show 15 / 15.',
        'Project path and Python executable must point to the current extracted release.',
        'Open http://127.0.0.1:8767/.',
    ])
    h(d,'2. Prepare one batch folder')
    numbered(d,[
        'Create one immediate child folder inside ~/SnapIMS-data/batches.',
        'Use the desired default batch name as the folder name.',
        'Put product photographs inside the folder in capture order.',
        'Photograph NEXT ITEM between tapes. The exact payload is CVHS1:ITEM:NEXT.',
        'Do not rename, recompress, or edit source photographs after Preview.',
    ])
    p(d,'Blank or “None” location means unassigned. Any reasonable free-text location is allowed, including Processing Table or Warehouse Wall.')
    h(d,'3. Preview and Commit')
    numbered(d,[
        'Open Import and select the folder.',
        'Confirm the derived batch name.',
        'Enter a location or leave it unassigned.',
        'Select Preview batch. The page remains responsive and progress survives refresh.',
        'Verify the interpreted Items and warning list.',
        'Use only the focused correction controls when necessary.',
        'Select Commit Import once grouping is correct. Repeated Commit is safe and idempotent.',
    ])
    image(d,'chromium-import-top.png','Import Preview overview with durable progress, measured scan counts, and the first interpreted Item.')
    image(d,'chromium-import-commit.png','Final Preview controls and Commit Import action.')
    h(d,'4. Recognition and Review')
    bullets(d,[
        'Recognition may use live AI only when configured. Blocked or failed results are not successful results.',
        'Review title, price, tags, condition, Rare, Physical Review, and evidence before approval.',
        'Manual edits remain attached to immutable Item IDs.',
    ])
    h(d,'5. Batch Editor')
    bullets(d,[
        'The editor uses bounded pages and does not render an entire 5,000-item batch.',
        'Select destination rows, focus a source cell, then choose Fill Down.',
        'Fill Down supports Title, Price, Tags, Discount, Description, Location, Review, and Rare.',
        'Tags Fill Down replaces the destination tag set exactly. Append Tags is a different action.',
        'Rows hidden by search/filter are automatically deselected.',
        'ArrowLeft and ArrowRight move through the Tags editor.',
    ])
    image(d,'chromium-batch-editor-controls.png','Desktop Batch Editor with controlled Tags, location, flags, paging, and selection-safe bulk actions.')
    image(d,'chromium-mobile-editor-top.png','At 390 pixels, the page body fits the viewport and the wide editor grid scrolls horizontally.')
    h(d,'6. Publish and finish')
    bullets(d,[
        'Publish readiness is calculated with aggregate database queries.',
        'Shopify simulation is explicit and limited to the current page.',
        'Live Shopify action requires valid credentials and confirmation.',
        'Stop SnapIMS before an offline database backup.',
    ])
    code(d,'snapims down')
    h(d,'End-of-batch check')
    numbered(d,[
        'Physical tape count matches Preview Item count.',
        'No NEXT photographs are attached to Items.',
        'Location is correct or intentionally unassigned.',
        'Titles, prices, tags, condition, flags, and quantity are verified.',
        'CSV and Shopify results were independently checked when used.',
    ])
    save(d,f'SnapIMS_v{VERSION}_Day-to-Day_Guide.docx')


def operator():
    d=setup('SnapIMS Operator Guide','Complete verified operator procedures')
    h(d,'1. Release identity and safety')
    p(d,'The active release is v0.13.2 with inventory schema 15. The installer pins the command and user service to the extracted release so an older source tree cannot silently start against a newer database.')
    table(d,['Identity','Expected'],[
        ('CLI version','0.13.2'),('Home version','0.13.2'),('Inventory schema','15 / 15'),('Default data','~/SnapIMS-data'),('Batch Home','~/SnapIMS-data/batches')])
    h(d,'2. Installation and startup')
    code(d,'unzip SnapIMS-v0.13.2-full-source.zip\ncd SnapIMS-v0.13.2\nchmod +x install_v0132.sh\n./install_v0132.sh\nexport PATH="$HOME/.local/bin:$PATH"\nsnapims status\nsnapims up')
    p(d,'Do not operate against an unexpected project path. Re-run the installer from the intended release if status shows a stale source path.')
    h(d,'3. Import protocol')
    p(d,'Folder beginning and folder end are automatic batch boundaries. NEXT ITEM is the only active daily QR command. START, END, LOCATION, RARE, REVIEW, and CONTINUE photographs are ignored legacy commands and produce one concise non-blocking warning.')
    table(d,['Situation','Interpretation'],[
        ('No NEXT card','Valid one-Item batch'),('Leading NEXT','Ignored with warning'),('Consecutive NEXT','Collapsed with warning'),('Trailing NEXT','Empty Item ignored with warning'),('Only NEXT cards','No product photographs'),('Unrelated consumer QR','Product photograph'),('Legacy SnapIMS QR','Ignored; does not control grouping')])
    h(d,'4. Import corrections')
    bullets(d,['Treat as Product Photo','Treat as NEXT ITEM','Ignore Photo','Split Item Before This Photo','Merge With Previous Item'])
    p(d,'Corrections affect only SnapIMS interpretation. They survive refresh, navigation, restart, Preview reopening, and final Commit. Original filenames, timestamps, EXIF, and pixels are not modified.')
    image(d,'chromium-import-top.png','Verified 100-item Preview overview: exact grouping, durable progress, and one decode/scan per source image.')
    image(d,'chromium-import-correction.png','Interpreted Item cards expose focused photo corrections without modifying source files.')
    image(d,'chromium-import-commit.png','End of Preview and Commit controls.')
    h(d,'5. Recognition')
    p(d,'Use configured providers only. Test/mock provider output is blocked from Publish unless deliberately replaced by manual review. A missing or failed provider is displayed as blocked or failed rather than silently accepted.')
    h(d,'6. Review')
    bullets(d,[
        'Inspect the correct Item photo and immutable Item ID.',
        'Confirm or change Title, Price, Tags, Discount, Condition, Quantity, and location.',
        'Use Rare and Physical Review as operator decisions.',
        'Review stored evidence and uncertainty; SnapIMS does not invent explanations for missing evidence.',
        'Stale two-tab edits return HTTP 409 and preserve the accepted record.',
    ])
    h(d,'7. Batch Editor')
    p(d,'The editor is paged at 20–200 rows. Search, filter, and sort are server-bounded, and the loaded page retains fast spreadsheet-style editing.')
    table(d,['Operation','Behavior'],[
        ('Fill Down Tags','Exact set/replace with validated immutable Tag IDs'),('Fill Down Location','Copies free text; blank/None becomes unassigned; history recorded'),('Fill Down Review/Rare','Copies true or false exactly'),('Append Tags','Adds tags without replacing existing tags'),('Hidden rows','Selections clear immediately'),('Keyboard','Horizontal movement crosses Tags; vertical movement keeps the field')])
    image(d,'chromium-batch-editor-controls.png','Desktop Batch Editor after the v0.13.2 consistency and paging fixes.')
    image(d,'chromium-mobile-editor-top.png','Mobile verification at 390 × 844.')
    h(d,'8. CSV')
    numbered(d,[
        'Export the intended batch CSV.',
        'Preserve Item ID values and header structure.',
        'Edit only supported values.',
        'Preview import differences before apply.',
        'Correct every validation error; failed staged imports remain all-or-nothing.',
        'Verify exported/imported row counts and an Item-ID sample.',
    ])
    h(d,'9. Publish')
    p(d,'The Publish page uses aggregate readiness counts. It does not dry-run all Items by default. Select simulation explicitly; simulation is paged. Live Shopify draft creation is guarded by valid configuration and explicit confirmation.')
    h(d,'10. Recovery')
    table(d,['Problem','Operator action'],[
        ('Wrong version/path','Run snapims status; re-run install_v0132.sh from the intended release.'),
        ('Interrupted scan','Reopen Import; resume the durable folder job.'),
        ('Duplicate tab/Commit','Open the existing job/result; do not create a new source folder copy.'),
        ('Migration issue','Stop SnapIMS; preserve data; restore only from copied pre-schema-v15 backup.'),
        ('API unavailable','Keep result blocked/failed; use Diagnostics; do not treat dummy evidence as production proof.'),
    ])
    h(d,'11. Shutdown and backup')
    code(d,'snapims down\ncp -a ~/SnapIMS-data ~/SnapIMS-data-backup-$(date +%Y%m%d-%H%M%S)')
    save(d,f'SnapIMS_v{VERSION}_Operator_Guide.docx')


def manual():
    d=setup('SnapIMS Complete Operating Manual','Operations, diagnostics, recovery, and release readiness')
    h(d,'1. Product model')
    p(d,'SnapIMS keeps immutable internal Batch and Item identities while allowing operator-visible batch names and locations to change. Source folders remain audit evidence and are never renamed by display-name edits.')
    table(d,['Layer','Purpose'],[
        ('Source photos','Immutable operator originals'),('Import cache','Restart-safe dimensions, readability, QR result, thumbnail, optional hash'),('Inventory database','Batches, Items, history, controlled Tags, checkpoints'),('Recognition','Durable provider attempts/results and evidence'),('Review/Batch Editor','Human decisions and audited edits'),('Publish','CSV and guarded Shopify draft workflow')])
    h(d,'2. Performance design')
    bullets(d,[
        'Folder list reads immediate child directories without opening every photograph.',
        'Preview runs as a background job and returns browser control immediately.',
        'Unchanged photographs are not decoded or QR-scanned again.',
        'Final preservation hashes are streamed at Commit and reused.',
        'Batch Editor and Publish use bounded database pages and aggregate summaries.',
        'SQLite connections close explicitly after every context.',
    ])
    table(d,['Measured scenario','v0.13.2 sandbox result'],[
        ('20-item / 42-photo cold Preview','479 ms'),('20-item warm Preview','18 ms, 0 decodes, 0 scans'),('100-item / 202-photo cold Preview','2,388 ms'),('100-item warm Preview','40 ms'),('5,000-item Batch Editor','430,721 bytes; 100 rows; 0.414 s relayed Chromium render'),('Full suite descriptors','3 baseline; 166 peak; 7 final; 0 final SQLite')])
    h(d,'3. Complete workflow')
    numbered(d,[
        'Start and verify release identity.',
        'Prepare one batch folder and NEXT separators.',
        'Preview and correct grouping.',
        'Commit immutable inventory records.',
        'Run Recognition or leave provider-blocked Items for manual work.',
        'Review evidence and commercial fields.',
        'Use paged Batch Editor for controlled mass edits.',
        'Verify CSV when used.',
        'Simulate and perform guarded Publish only when ready.',
        'Stop and back up.',
    ])
    image(d,'chromium-import-top.png','Import Preview overview from the v0.13.2 browser pass.')
    image(d,'chromium-import-commit.png','Import Commit controls from the v0.13.2 browser pass.')
    image(d,'chromium-batch-editor-controls.png','Batch Editor evidence from the v0.13.2 browser pass.')
    h(d,'4. Audit and conflict rules')
    bullets(d,[
        'Every accepted edit remains tied to the immutable Item ID.',
        'Location changes create location-history events.',
        'Bulk operations are atomic and checkpointed where required.',
        'Stale revisions are rejected rather than overwriting newer work.',
        'Duplicate Import Preview/Commit requests resolve to one durable job/result.',
        'A mixed application/schema version is refused before serving normal work.',
    ])
    h(d,'5. Diagnostics')
    code(d,'snapims doctor\nsnapims status\nsnapims logs')
    p(d,'Status reports project path, Python executable, data directory, PID, version, and schema. External-service failures do not imply inventory database failure; read each component separately.')
    h(d,'6. Security')
    bullets(d,[
        'Configure authentication before public exposure.',
        'Keep secrets outside source control and verify owner-only secret-file permissions.',
        'Use CSRF-protected state changes and revoke sessions after credential changes.',
        'Treat Cloudflare, Guacamole, xrdp, and remote-shell exposure as separate deployment controls.',
    ])
    h(d,'7. Production acceptance')
    p(d,'v0.13.2 is pilot-ready, not v1.0.0. Before 1.0, complete the real 20-tape Pixel pilot, live AI, live Shopify draft, real CSV round trip, target-machine restart durability, direct Firefox/Chromium walkthrough, guide verification, security/deployment review, backup/restore drill, and zero-blocker review.')
    h(d,'8. Skipped sandbox capabilities')
    table(d,['Capability','Reason'],[
        ('Live OpenAI','No production API key'),('Live Shopify','No store/token'),('Native Firefox','Runtime unavailable'),('Direct localhost Chromium','Managed URLBlocklist'),('2012 Mac mini','Hardware unavailable'),('Cloudflare/Guacamole','Owner infrastructure unavailable'),('HEIC','pillow_heif unavailable')])
    h(d,'9. Evidence locations')
    bullets(d,[
        'release-evidence/v0.13.2/',
        'browser-evidence/v0.13.2/',
        'TEST_RESULTS.md',
        'PERFORMANCE_REPORT.md',
        'ISSUE_TRACEABILITY_V0132.md',
        'V1_PRODUCTION_ACCEPTANCE_GAPS.md',
    ])
    save(d,f'SnapIMS_v{VERSION}_Complete_Operating_Manual.docx')

if __name__=='__main__':
    day_to_day(); operator(); manual()
